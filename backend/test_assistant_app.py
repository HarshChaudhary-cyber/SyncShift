"""Regression tests for direct schedule actions and document-to-calendar imports."""
import asyncio
import io
from datetime import date, time, timedelta

import pytest
from fastapi import UploadFile, HTTPException
from starlette.requests import Request
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.dependencies import CurrentUser
from app.models.user import User
from app.models.time_block import TimeBlock
from app.models.block_override import BlockOverride
from app.services.assistant_app import app_help, is_today_work_move, move_today_work
from app.services.file_extractor import extract_text
from app.routers.import_file import preview_file_timetable, confirm_file_import
from app.schemas.import_ics import IcsConfirmRequest
from app.store import get_occurrences_for_range


@pytest.fixture
def demo():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add_all([User(id=1, email="one@example.com"), User(id=2, email="two@example.com")])
        db.commit()
        yield db, CurrentUser(1, "one@example.com", timezone="Asia/Kolkata", minimum_transition_minutes=0)
    engine.dispose()


def block(db, user_id=1, day=1, kind="shift", start=time(10), end=time(12), recurring=True, specific=None):
    b = TimeBlock(user_id=user_id, title="Test work", type=kind, day_of_week=day,
        start_time=start, end_time=end, duration_minutes=120, is_recurring=recurring,
        specific_date=specific)
    db.add(b); db.commit()
    return b


def test_direct_move_is_one_occurrence_and_repeat_is_noop(demo):
    db, user = demo
    today = date(2026, 9, 28)
    own = block(db)
    other = block(db, user_id=2)
    text, changed = move_today_work(db, user, today)
    assert changed and "2026-09-29" in text
    assert not get_occurrences_for_range(1, today, today, db)
    assert len(get_occurrences_for_range(1, today+timedelta(days=1), today+timedelta(days=1), db)) == 1
    assert len(get_occurrences_for_range(1, today+timedelta(days=7), today+timedelta(days=7), db)) == 1
    assert db.get(TimeBlock, own.id).day_of_week == 1
    assert db.get(TimeBlock, other.id).day_of_week == 1
    assert move_today_work(db, user, today)[1] is False
    assert db.query(BlockOverride).count() == 1


def test_conflict_makes_no_partial_changes(demo):
    db, user = demo
    block(db)
    block(db, start=time(15), end=time(17))
    block(db, day=2, kind="class", start=time(16), end=time(18))
    text, changed = move_today_work(db, user, date(2026, 9, 28))
    assert not changed and "haven't moved" in text
    assert db.query(BlockOverride).count() == 0


def test_move_single_date_and_ignore_other_users_conflict(demo):
    db, user = demo
    today = date(2026, 9, 28)
    own = block(db, recurring=False, specific=today)
    block(db, user_id=2, day=2, kind="class")
    assert move_today_work(db, user, today)[1]
    assert own.specific_date == today+timedelta(days=1)


def test_help_and_explicit_command_boundary():
    assert "Smart Planner" in app_help("How do I use this whole app?", "student")
    assert is_today_work_move("Can you move my today work to tommorrow?")
    assert not is_today_work_move("Do not move my today work to tomorrow")
    assert not is_today_work_move("How do I move my today work to tomorrow?")
    assert not is_today_work_move("Move my today work to tomorrow at 6pm")


def test_sunday_move_respects_next_weeks_limit(demo):
    db, user = demo
    user.weekly_work_hour_limit = 3
    block(db, day=0)
    block(db, day=2)
    message, changed = move_today_work(db, user, date(2026, 9, 27))
    assert not changed and "work-hour limit" in message
    assert db.query(BlockOverride).count() == 0


def test_docx_table_extraction():
    from docx import Document
    document = Document()
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Monday", "09:00-10:00", "Physics"]):
        cell.text = value
    data = io.BytesIO(); document.save(data)
    text, is_ocr = extract_text(data.getvalue(), "timetable.docx")
    assert "Monday" in text and "Physics" in text and not is_ocr


@pytest.mark.parametrize("extension", ["png", "pdf"])
def test_local_scanned_timetable_ocr(extension):
    import shutil
    if not shutil.which("tesseract"):
        pytest.skip("OCR runtime is installed in the Docker image")
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1200, 180), "white")
    draw = ImageDraw.Draw(image)
    draw.text((35, 50), "Monday 09:00-10:00 Physics Room 4",
              font=ImageFont.truetype("DejaVuSans.ttf", 38), fill="black")
    data = io.BytesIO(); image.save(data, format=extension.upper())
    text, is_ocr = extract_text(data.getvalue(), f"scan.{extension}")
    assert "Physics" in text and "09:00" in text and is_ocr


def test_excel_csv_preview_and_import_day_alignment(demo):
    import openpyxl
    db, user = demo
    workbook = openpyxl.Workbook()
    workbook.active.append(["Title", "Day", "Start", "End", "Room"])
    workbook.active.append(["CS101", "Monday", time(9), time(10), "Room 4"])
    buf = io.BytesIO(); workbook.save(buf)
    response = asyncio.run(preview_file_timetable(
        UploadFile(filename="schedule.xlsx", file=io.BytesIO(buf.getvalue())), user, db, None)).data
    assert len(response.preview) == 1
    row = response.preview[0]
    assert row.day_of_week == 0 and row.start_time == "09:00"
    payload = IcsConfirmRequest(preview_blocks=[row.model_dump()])
    request = Request({"type": "http", "headers": []})
    assert confirm_file_import(payload, request, user, db).data.created_count == 1
    assert db.query(TimeBlock).one().day_of_week == 1
    assert confirm_file_import(payload, request, user, db).data.created_count == 0
    text, _ = extract_text(b'Title,Day,Start,End\n"Math, advanced",Tuesday,13:00,14:00', "schedule.csv")
    assert "Tuesday 13:00 - 14:00 Math, advanced" in text


def test_import_invalid_time_rolls_back_batch(demo):
    db, user = demo
    payload = IcsConfirmRequest(preview_blocks=[
        dict(title="Valid", day_of_week=0, start_time="09:00", end_time="10:00"),
        dict(title="Invalid", day_of_week=0, start_time="25:00", end_time="26:00"),
    ])
    with pytest.raises(HTTPException) as exc:
        confirm_file_import(payload, Request({"type": "http", "headers": []}), user, db)
    assert exc.value.status_code == 422
    assert db.query(TimeBlock).count() == 0


def test_published_class_blocks_automatic_import(demo, monkeypatch):
    from app.schemas.block import BlockOut
    db, user = demo
    official = BlockOut(id=-5, type="class", title="Official lecture", day_of_week=1,
                        start_time="09:00", end_time="10:00")
    monkeypatch.setattr("app.store.get_occurrences_for_range", lambda *args, **kwargs: [official])
    result = asyncio.run(preview_file_timetable(UploadFile(filename="schedule.csv",
        file=io.BytesIO(b"Day,Start,End,Title\nMonday,09:00,10:00,Math")), user, db, None)).data
    assert result.preview[0].has_conflict
    assert result.preview[0].status == "conflict"
